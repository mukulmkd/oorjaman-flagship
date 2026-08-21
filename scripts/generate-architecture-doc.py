#!/usr/bin/env python3
"""
Generate OorjaMan complete architecture documentation (Word) in project-docs/.

Run: .venv-docgen/bin/python scripts/generate-architecture-doc.py
     or: npm run docs:architecture

Regenerate after major schema, app, or deployment changes.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SCRIPTS_DIR = Path(__file__).resolve().parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from architecture_diagram_boxes import render_master_diagram_section  # noqa: E402

REPO_ROOT = SCRIPTS_DIR.parent
OUTPUT = REPO_ROOT / "project-docs" / "OorjaMan-Architecture.docx"

ADMIN_UAT = "https://oorjaman-admin.vercel.app"
VENDOR_UAT = "https://oorjaman-vendor.vercel.app"
SUPPORT_UAT = "https://oorjaman-support.vercel.app"


def shade_cell(cell, fill: str = "E8F0FE") -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(el)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def keep_table_on_one_page(table) -> None:
    for row in table.rows:
        prevent_row_split(row)


def set_cell_padding(cell, top=80, bottom=80, left=100, right=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        shade_cell(table.rows[0].cells[i])
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.rows[r + 1].cells[c].text = val
    doc.add_paragraph()


def add_diagram(
    doc: Document,
    title: str,
    lines: list[str],
    font_size: float = 8.5,
    *,
    own_page: bool = True,
) -> None:
    """Monospace diagram inside a single-cell table that must not split across pages."""
    if own_page:
        doc.add_page_break()
    if title:
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x5F, 0x4A)
    frame = doc.add_table(rows=1, cols=1)
    frame.style = "Table Grid"
    cell = frame.rows[0].cells[0]
    shade_cell(cell, "FAFAFA")
    set_cell_padding(cell, top=60, bottom=60, left=80, right=80)
    cell.text = ""
    for line in lines:
        para = cell.add_paragraph()
        run = para.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(font_size)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
    keep_table_on_one_page(frame)
    doc.add_paragraph()


def add_diagram_page_block(
    doc: Document,
    heading: str,
    lines: list[str],
    *,
    title: str = "",
    level: int = 2,
    intro: str = "",
    font_size: float = 8.5,
) -> None:
    """Heading + diagram kept together on one page (page break before the block)."""
    doc.add_page_break()
    doc.add_heading(heading, level=level)
    if intro:
        doc.add_paragraph(intro)
    add_diagram(doc, title, lines, font_size=font_size, own_page=False)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    return doc


def add_cover(doc: Document) -> None:
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("OorjaMan Flagship\nComplete Architecture Documentation")
    run.bold = True
    run.font.size = Pt(26)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Monorepo · Apps · Database · Integrations · {date.today().strftime('%d %B %Y')}")

    doc.add_paragraph()
    doc.add_paragraph(
        "This document describes how all OorjaMan applications interact, the shared Supabase "
        "backend, database schema, security model, and deployment topology. It is the canonical "
        "architecture reference for engineers, architects, and technical stakeholders."
    )
    doc.add_page_break()


def add_toc_placeholder(doc: Document) -> None:
    doc.add_heading("Table of contents", level=1)
    sections = [
        "1. Executive summary",
        "2. Master system architecture diagram (full detail)",
        "3. Monorepo structure",
        "4. Applications",
        "5. Shared packages",
        "6. Application interactions",
        "7. Authentication and roles",
        "8. Core business flows",
        "9. Notifications and realtime",
        "10. Supabase platform",
        "11. Database architecture",
        "12. Row-level security (RLS)",
        "13. Edge functions",
        "14. Storage buckets",
        "15. Deployment architecture",
        "16. Environments and configuration",
        "17. External services",
        "18. Security overview",
        "19. Appendix",
    ]
    add_bullets(doc, sections)
    doc.add_paragraph(
        "Tip: In Microsoft Word, use References → Table of Contents to auto-generate a TOC "
        "from headings after opening this file."
    )
    doc.add_page_break()


def add_executive_summary(doc: Document) -> None:
    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(
        "OorjaMan is a solar rooftop care platform delivered as an npm workspaces monorepo. "
        "Six client applications (two mobile, three operations portals, one marketing site) "
        "share a single TypeScript domain layer (@oorjaman/api) and coordinate exclusively "
        "through Supabase — there is no custom REST API server."
    )
    add_table(
        doc,
        ["Principle", "Description"],
        [
            ["Backend-as-a-Service", "Postgres + Auth + Realtime + Storage + Edge Functions on Supabase"],
            ["Shared domain logic", "All business rules in packages/api; clients are thin UI shells"],
            ["Role-separated clients", "customer, technician, vendor, admin, support each have dedicated apps"],
            ["Two Supabase projects", "OorjaMan UAT (dev/staging) and OorjaMan Prod (live data isolation)"],
            ["Payments (current)", "Simulated dummy gateway; Razorpay/Stripe not yet integrated"],
        ],
    )


def add_master_architecture_diagram(doc: Document) -> None:
    render_master_diagram_section(doc, add_table)


def add_monorepo(doc: Document) -> None:
    doc.add_heading("3. Monorepo structure", level=1)
    add_diagram(
        doc,
        "Diagram 3.1 — Repository layout",
        [
            "oorjaman-flagship/",
            "├── apps/",
            "│   ├── customer-app/      Expo mobile — homeowners",
            "│   ├── technician-app/    Expo mobile — field partners (OorjaMan Partner)",
            "│   ├── admin-web/         Vite SPA — platform operations",
            "│   ├── vendor-web/        Vite SPA — solar cleaning partners",
            "│   ├── support-web/       Vite SPA — customer & technician support desk",
            "│   └── oorjaman-web/      Next.js — public marketing site",
            "├── packages/",
            "│   ├── api/               Domain logic + Supabase types",
            "│   ├── ui/                React Native shared components",
            "│   ├── web-ui/            Vite portal shell + shared web components",
            "│   ├── config/ · utils/   Design tokens + pure helpers",
            "│   ├── mobile-deps/       Shared Expo/RN dependency bundle",
            "│   ├── mobile-config/     Expo plugins, Metro, native scripts",
            "│   ├── portal-deps/       Shared portal dependency bundle",
            "│   └── vite-portal-config/ Shared Vite/ESLint/TS config",
            "├── supabase/",
            "│   ├── migrations/        Schema source of truth (109+ files)",
            "│   └── functions/         Edge functions (push, notifications, intake)",
            "├── scripts/               DB push, seed, brand sync, doc generators",
            "└── project-docs/          Runbooks + generated Word specs",
        ],
        own_page=False,
    )


def add_applications(doc: Document) -> None:
    doc.add_heading("4. Applications", level=1)
    add_table(
        doc,
        ["App", "Audience", "Stack", "Local port", "Deployment"],
        [
            ["customer-app", "Homeowners / businesses", "Expo 56, RN 0.85, Expo Router", "Metro 8081", "EAS → App Store / Play"],
            ["technician-app", "Field technicians", "Expo 56, RN 0.85", "Metro 8081", "EAS → App Store / Play"],
            ["admin-web", "Platform operators", "React 19, Vite 8, React Router 7", "5173", "Vercel UAT → GoDaddy prod"],
            ["vendor-web", "Approved vendors", "React 19, Vite 8", "5174", "Vercel UAT → GoDaddy prod"],
            ["support-web", "Support desk staff", "React 19, Vite 8", "5175", "Vercel UAT → GoDaddy prod"],
            ["oorjaman-web", "Public visitors", "Next.js 15 (static export)", "3000", "GoDaddy (oorjaman.com)"],
        ],
    )

    doc.add_heading("4.1 Mobile apps (PROD vs UAT)", level=2)
    add_table(
        doc,
        ["", "Production", "UAT"],
        [
            ["Customer display name", "OorjaMan", "OorjaMan (UAT)"],
            ["Partner display name", "OorjaMan Partner", "OorjaMan Partner (UAT)"],
            ["iOS bundle (customer)", "com.oorjaman.customer", "com.oorjaman.customer.uat"],
            ["Deep link scheme", "oorjaman-customer", "oorjaman-customer-uat"],
            ["Env file (embedded builds)", ".env.production.local", ".env.uat.local"],
        ],
    )

    doc.add_heading("4.2 Web portals (UAT URLs)", level=2)
    add_table(
        doc,
        ["Portal", "UAT URL", "Production target"],
        [
            ["Admin", ADMIN_UAT, "https://admin.oorjaman.com"],
            ["Vendor", VENDOR_UAT, "https://vendor.oorjaman.com"],
            ["Support", SUPPORT_UAT, "https://support.oorjaman.com"],
            ["Marketing", "dev-oorjaman.oorjaman.com (planned)", "https://oorjaman.com"],
        ],
    )


def add_packages(doc: Document) -> None:
    doc.add_heading("5. Shared packages", level=1)
    add_table(
        doc,
        ["Package", "Responsibility"],
        [
            ["@oorjaman/api", "Supabase client factories, typed models, all domain APIs (bookings, AMC, payments, finance, support, notifications, analytics)"],
            ["@oorjaman/ui", "React Native: screens, modals, OTP, offline gate, mobile notification handlers"],
            ["@oorjaman/web-ui", "React web: Supabase provider, session gates, portal login, notification bell, document viewer"],
            ["@oorjaman/config", "Design tokens, brand assets, deploy-env helpers"],
            ["@oorjaman/utils", "IST dates, booking slots, brand-print layouts"],
            ["@oorjaman/mobile-deps", "Shared Expo/RN runtime versions for both mobile apps"],
            ["@oorjaman/mobile-config", "Expo config plugins, Metro, prebuild/rebuild scripts"],
            ["@oorjaman/portal-deps", "Aggregates portal runtime dependencies"],
            ["@oorjaman/vite-portal-config", "createPortalViteConfig() with monorepo path aliases"],
        ],
    )
    doc.add_page_break()
    add_diagram(
        doc,
        "Diagram 5.1 — Dependency flow",
        [
            "  customer-app ──┐",
            "  technician-app ├──► mobile-deps ──► api ──► Supabase",
            "                 │         │",
            "                 └────────► mobile-config",
            "",
            "  admin/vendor/support-web ──► portal-deps ──► web-ui ──► api",
            "  oorjaman-web ─────────────────────────────► config, utils",
        ],
        own_page=False,
    )


def add_interactions(doc: Document) -> None:
    doc.add_heading("6. Application interactions", level=1)
    doc.add_paragraph(
        "See Section 2 (Master system architecture diagram) for the complete interaction map. "
        "Applications never call each other directly — coordination happens through shared "
        "database rows, Supabase Realtime channels, notification_events, and edge functions."
    )

    add_diagram(
        doc,
        "Diagram 6.1 — Communication mechanisms",
        [
            "┌─────────────┐     postgres_changes      ┌─────────────┐",
            "│ customer-app│◄────────────────────────►│  bookings   │",
            "└─────────────┘                           │subscriptions│",
            "┌─────────────┐     postgres_changes      │notification_│",
            "│technician-app│◄────────────────────────►│   events    │",
            "└─────────────┘                           │support_msgs │",
            "┌─────────────┐     postgres_changes      │vendor_settl.│",
            "│ vendor-web  │◄────────────────────────►└─────────────┘",
            "└─────────────┘                                    ▲",
            "┌─────────────┐     postgres_changes                 │",
            "│ admin-web   │◄─────────────────────────────────────┘",
            "└─────────────┘",
            "",
            "Mobile push: support message → push_outbox → edge function → Expo Push API",
            "External notify: notification_events → process-notification-events → email/SMS/WhatsApp (demo)",
        ],
    )

    doc.add_heading("6.1 Who reads/writes what", level=2)
    add_table(
        doc,
        ["Data domain", "customer-app", "technician-app", "vendor-web", "admin-web", "support-web"],
        [
            ["bookings", "Create, track, cancel", "Execute visit workflow", "Accept, assign, cancel", "Monitor, ops intervene", "Read for context"],
            ["subscriptions / AMC", "Subscribe, book visits", "—", "View assigned AMC", "Assign vendor, pricing", "Read context"],
            ["payments", "Pay (dummy)", "—", "View related", "Finance console", "—"],
            ["vendor_settlements", "—", "—", "View own payouts", "Approve / settle", "—"],
            ["support_conversations", "Create, chat", "Create, chat", "—", "—", "Full desk"],
            ["notification_events", "—", "—", "Vendor inbox", "Admin inbox", "—"],
            ["pricing catalog", "Read (quotes)", "—", "—", "CRUD", "—"],
            ["vendor intake", "—", "—", "Signup wizard", "Approve", "—"],
        ],
    )


def add_auth(doc: Document) -> None:
    doc.add_heading("7. Authentication and roles", level=1)
    add_diagram(
        doc,
        "Diagram 7.1 — Auth pipeline",
        [
            "  User enters phone/email OTP",
            "           │",
            "           ▼",
            "  Supabase Auth (auth.users) ──trigger──► public.users (role, profile)",
            "           │",
            "           ▼",
            "  Role extension row: customers | vendors | technicians | support_agents",
            "           │",
            "           ▼",
            "  App-specific gate: RequireSession / wrong-role screen / approval check",
        ],
        own_page=False,
    )
    add_table(
        doc,
        ["Role (public.users.role)", "App", "Gate"],
        [
            ["customer", "customer-app", "Onboarding → main tabs; wrong-role for other apps"],
            ["technician", "technician-app", "Vendor link + onboarding; must be technician role"],
            ["vendor", "vendor-web", "RequireVendorRole + RequireApprovedVendor for dashboard"],
            ["admin", "admin-web", "RequireAdminRole only"],
            ["support", "support-web", "RequireSupportDeskRole (admin or support)"],
            ["—", "oorjaman-web", "No authentication (public marketing)"],
        ],
    )
    doc.add_paragraph(
        "UAT uses dummy auth (OTP 123456) when USE_DUMMY_AUTH=true. Production uses real "
        "Supabase Phone/Email providers."
    )


def add_flows(doc: Document) -> None:
    doc.add_heading("8. Core business flows", level=1)

    add_diagram_page_block(
        doc,
        "8.1 One-time booking (happy path)",
        [
            "Customer          Supabase           Vendor           Technician        Admin",
            "   │                  │                  │                  │              │",
            "   │ create booking   │                  │                  │              │",
            "   │ (pending_payment)│                  │                  │              │",
            "   ├─────────────────►│                  │                  │              │",
            "   │ dummy payment    │                  │                  │              │",
            "   ├─────────────────►│ confirmed        │                  │              │",
            "   │                  ├─notification────►│ assigned         │              │",
            "   │                  ├─notification──────────────────────────────────────────►│",
            "   │                  │                  │ accept (1h win)  │              │",
            "   │                  │◄─────────────────┤ assign tech      │              │",
            "   │                  ├─realtime────────────────────────────►│ job appears  │",
            "   │◄─realtime────────┤ status updates   │                  │ in_progress  │",
            "   │ track technician │                  │                  ├─completed───►│",
            "   │                  │                  │                  │ settlement   │",
        ],
        title="Diagram 8.1 — Booking sequence",
    )
    doc.add_paragraph("Booking status progression:")
    add_bullets(
        doc,
        [
            "pending_payment → confirmed → vendor_acknowledged → accepted → in_progress → completed",
            "Alternative terminal state: cancelled (customer, vendor, or admin)",
            "Vendor has 1 hour to accept/reject; overdue handled by pg_cron + scan-vendor-response-overdue",
            "Marketplace fallback: admin floats job when no preferred vendor available",
        ],
    )

    add_diagram_page_block(
        doc,
        "8.2 AMC / subscription flow",
        [
            "Customer subscribes (subscriptions)",
            "        │",
            "        ▼",
            "ensureAmcWalletForSubscription (pending_funding)",
            "        │",
            "        ▼",
            "Payment success → fund_amc_wallet_from_payment RPC",
            "        │",
            "        ▼",
            "Admin assigns assigned_vendor_id",
            "        │",
            "        ▼",
            "syncAmcVisitSlotsForSubscription → customer schedules slots",
            "        │",
            "        ▼",
            "scheduleAmcVisitSlot → booking (subscription_id set)",
            "        │",
            "        ▼",
            "Visit complete → release_amc_wallet_visit_payout (not standard visit settlement)",
        ],
        title="Diagram 8.2 — AMC lifecycle",
    )

    doc.add_heading("8.3 Vendor onboarding", level=2)
    add_bullets(
        doc,
        [
            "vendor-web /signup → vendor_registration_intake + documents (vendor-intake bucket)",
            "admin-web Vendor Approval reviews intake",
            "approve-vendor-intake edge function → auth.users + vendors row + doc copy",
            "Vendor logs in → pending until approval_status = approved",
            "Vendor invites technicians → technicians row → technician-app onboarding",
        ],
    )

    doc.add_heading("8.4 Support chat", level=2)
    add_bullets(
        doc,
        [
            "Customer or technician opens support_conversations from mobile apps",
            "support-web inbox + SupportChatDock; Realtime on support_messages",
            "When app is killed: customer_push_outbox / technician_push_outbox → Expo push edge functions",
        ],
    )


def add_notifications(doc: Document) -> None:
    doc.add_heading("9. Notifications and realtime", level=1)
    add_diagram(
        doc,
        "Diagram 9.1 — Notification architecture (dual track)",
        [
            "Track A — In-app (Postgres + Realtime)",
            "  API inserts notification_events (channels: in_app, …)",
            "       → admin-web / vendor-web subscribe via postgres_changes",
            "       → optional process-notification-events for email/SMS/WhatsApp queue",
            "",
            "Track B — Mobile push (Expo)",
            "  Support message → enqueue customer_push_outbox / technician_push_outbox",
            "       → send-customer-expo-push / send-technician-expo-push",
            "       → Expo Push API → device",
        ],
        own_page=False,
    )
    add_table(
        doc,
        ["Realtime table", "Subscribers"],
        [
            ["bookings", "customer-app, technician-app, vendor-web"],
            ["subscriptions", "customer-app (AMC partner assigned)"],
            ["notification_events", "admin-web, vendor-web (in-app bell)"],
            ["support_messages, support_conversations", "customer-app, technician-app, support-web"],
            ["vendor_settlements", "vendor-web"],
            ["customer_site_activity_events", "customer-app site timeline"],
            ["technician_activity_events", "technician-app activity feed"],
        ],
    )


def add_supabase_platform(doc: Document) -> None:
    doc.add_heading("10. Supabase platform", level=1)
    add_table(
        doc,
        ["Component", "Purpose in OorjaMan"],
        [
            ["Postgres", "All business data; RLS enforces per-role access"],
            ["Auth", "Phone/email OTP; JWT for client requests"],
            ["PostgREST", "Implicit REST API used by supabase-js (no custom backend)"],
            ["Realtime", "Live booking updates, notification bells, support chat"],
            ["Storage", "Site photos, job evidence, vendor/technician documents"],
            ["Edge Functions", "Push dispatch, notification processing, vendor intake approval, cron hooks"],
            ["pg_cron / pg_net", "Scheduled vendor overdue scan, push function invocation"],
        ],
    )
    doc.add_paragraph(
        "Schema evolves via timestamped SQL files in supabase/migrations/. Deploy with "
        "npm run db:push on each Supabase project (UAT first, then Prod). Reference snapshots: "
        "supabase/schema.sql, supabase/policies.sql."
    )


def add_database(doc: Document) -> None:
    doc.add_heading("11. Database architecture", level=1)

    doc.add_heading("11.1 Identity layer", level=2)
    add_table(
        doc,
        ["Table", "Purpose", "Key relationships"],
        [
            ["auth.users", "Supabase credentials (managed by Auth)", "id = public.users.id"],
            ["public.users", "App profile: role, contact, verification", "1:1 with auth.users"],
            ["customers", "Customer extension (site, solar sizing)", "user_id → users"],
            ["vendors", "Partner org (approval, regions, docs)", "user_id → users"],
            ["technicians", "Field crew profile", "user_id → users, vendor_id → vendors"],
            ["vendor_registration_intake", "Pre-auth signup drafts", "→ user/vendor on approval"],
            ["support_agents", "Support desk staff", "user_id → users (role support)"],
        ],
    )

    doc.add_heading("11.2 Operations layer", level=2)
    add_table(
        doc,
        ["Table", "Purpose"],
        [
            ["bookings", "Visit jobs (one-time + AMC); status machine"],
            ["subscriptions", "AMC contracts per service address"],
            ["subscription_visit_slots", "Planned AMC visit windows"],
            ["payments", "INR paise ledger (dummy gateway today)"],
            ["job_reports", "Completion evidence, ratings"],
            ["technician_locations", "GPS trail during active jobs"],
            ["vendor_slot_availability", "Partner capacity calendar"],
            ["vendor_technician_invites", "Team onboarding invites"],
            ["platform_settings", "Singleton ops config (fees, default vendor, support desk)"],
        ],
    )

    doc.add_heading("11.3 Finance layer", level=2)
    add_table(
        doc,
        ["Table", "Purpose"],
        [
            ["amc_wallets", "Customer-funded escrow per subscription"],
            ["amc_wallet_entries", "Ledger: fund, visit_release, platform_fee, refund"],
            ["vendor_settlements", "visit_payout or cancellation_penalty per booking"],
            ["vendor_deferred_penalties", "Penalties when no payout to net against"],
            ["customer_oorjaman_credit_grants / _redemptions", "Customer credits (e.g. late vendor cancel)"],
        ],
    )

    doc.add_heading("11.4 Pricing catalog", level=2)
    add_bullets(
        doc,
        [
            "pricing_tiers, pricing_city_tiers, pricing_rules",
            "service_capacity_tiers, pricing_one_time_rates, pricing_amc_plans",
            "pricing_national_default_audit, pricing_catalog_audit",
        ],
    )

    doc.add_heading("11.5 Notifications", level=2)
    add_table(
        doc,
        ["Table", "Purpose"],
        [
            ["notification_events", "Queue + in-app inbox (admin/vendor audiences)"],
            ["notification_templates", "Per event_type × channel copy"],
            ["notification_channel_settings", "Demo vs live channel toggles"],
            ["customer_push_tokens / technician_push_tokens", "Expo device registry"],
            ["customer_push_outbox / technician_push_outbox", "Push delivery queue (service_role only)"],
        ],
    )

    doc.add_heading("11.6 Support desk", level=2)
    add_bullets(
        doc,
        [
            "support_conversations, support_messages, support_message_attachments",
            "support_conversation_events, support_agents, support_macros",
        ],
    )

    doc.add_heading("11.7 Key enums", level=2)
    add_table(
        doc,
        ["Enum", "Values"],
        [
            ["user_role", "customer, vendor, technician, admin, support"],
            ["booking_status", "pending_payment, confirmed, vendor_acknowledged, accepted, in_progress, completed, cancelled"],
            ["vendor_approval_status", "pending, under_review, approved, rejected, suspended"],
            ["subscription_status", "trialing, active, paused, cancelled, expired, past_due"],
            ["payment_status", "pending, success, failed"],
            ["vendor_settlement_kind", "visit_payout, cancellation_penalty"],
        ],
    )

    doc.add_heading("11.8 Analytics views", level=2)
    add_bullets(
        doc,
        [
            "booking_stats, revenue_stats, recognized_revenue_stats, finance_dashboard_stats",
            "vendor_stats, technician_stats, subscription_stats, bookings_created_daily",
            "ops_booking_exceptions (admin ops desk exception queue)",
        ],
    )

    add_diagram_page_block(
        doc,
        "Entity relationship (core tables)",
        [
            "  users ─────┬──── customers ────┬──── bookings ◄──── technicians",
            "             │                     │         ▲              ▲",
            "             ├──── vendors ────────┴─────────┘              │",
            "             │         │                                        │",
            "             └──── technicians (vendor_id) ─────────────────────┘",
            "",
            "  customers ─── subscriptions ─── amc_wallets ─── amc_wallet_entries",
            "                    │                │",
            "                    └── subscription_visit_slots",
            "                    └── bookings (AMC visits)",
            "",
            "  bookings ─── job_reports · vendor_settlements · notification_events",
            "  customers ─── payments",
        ],
        title="Diagram 11.1 — Entity relationship (core)",
        level=2,
    )


def add_rls(doc: Document) -> None:
    doc.add_heading("12. Row-level security (RLS)", level=1)
    add_diagram(
        doc,
        "Diagram 12.1 — RLS helper functions",
        [
            "is_admin()              → users.role = 'admin'",
            "my_customer_id()        → customers.id for auth.uid()",
            "my_vendor_id()          → vendors.id for auth.uid()",
            "my_technician_id()      → technicians.id for auth.uid()",
            "is_approved_vendor_user() → vendor with approval_status = approved",
            "is_support_desk_user()  → admin or support agent",
        ],
        own_page=False,
    )
    add_table(
        doc,
        ["Domain", "Read access", "Write access"],
        [
            ["customers", "Self, admin, vendor if shared booking", "Self (onboarding fields)"],
            ["vendors", "Self, admin, team technician; customers see approved only", "Self; approval fields trigger-guarded"],
            ["bookings", "Customer own, assigned vendor/tech, support desk, admin", "Per role workflow (create, accept, execute)"],
            ["subscriptions", "Customer, admin, assigned vendor", "Customer/admin create; admin assigns AMC vendor"],
            ["payments", "Customer own; admin; vendor if linked booking", "Customer pending rows"],
            ["notification_events", "Admin all; vendor own inbox", "Authenticated insert; admin/vendor mark read"],
            ["support_*", "Participant + desk users", "Desk + participants per conversation"],
            ["push outbox", "No client access", "Triggers / edge functions only"],
        ],
    )
    doc.add_paragraph(
        "Marketplace: approved vendors can SELECT/UPDATE floated confirmed bookings "
        "(metadata.marketplace.mode=default_vendor, floated=true) before vendor_id is assigned."
    )


def add_edge_functions(doc: Document) -> None:
    doc.add_heading("13. Edge functions", level=1)
    add_table(
        doc,
        ["Function", "Purpose", "Auth / trigger", "Called by"],
        [
            ["process-notification-events", "Drain notification_events queue; render templates; demo email/SMS/WhatsApp", "Admin JWT", "adminProcessNotificationQueue()"],
            ["approve-vendor-intake", "Approve intake → create auth user + vendor + copy docs", "Admin JWT", "adminApproveVendorRegistrationIntake()"],
            ["scan-vendor-response-overdue", "RPC notify_overdue_vendor_responses_batch", "CRON_DISPATCH_SECRET", "pg_cron every 5 min"],
            ["send-customer-expo-push", "Drain customer_push_outbox → Expo API", "PUSH_DISPATCH_SECRET", "pg_net / cron"],
            ["send-technician-expo-push", "Drain technician_push_outbox → Expo API", "PUSH_DISPATCH_SECRET", "pg_net / cron"],
        ],
    )


def add_storage(doc: Document) -> None:
    doc.add_heading("14. Storage buckets", level=1)
    add_table(
        doc,
        ["Bucket", "Public", "Purpose"],
        [
            ["customer-site-photos", "No", "Customer rooftop/site images (address-scoped)"],
            ["job-photos", "Yes (unguessable paths)", "Technician visit evidence"],
            ["vendor-documents", "No", "Approved partner compliance docs"],
            ["technician-documents", "No", "Technician KYC / certifications"],
            ["vendor-intake", "No", "Pre-approval signup uploads"],
            ["support-attachments", "No", "Support desk file attachments"],
        ],
    )


def add_deployment(doc: Document) -> None:
    doc.add_heading("15. Deployment architecture", level=1)
    add_diagram(
        doc,
        "Diagram 15.1 — Deployment topology",
        [
            "                         ┌── OorjaMan UAT Supabase ──┐",
            "  Local dev ──────────────┤                           ├── Vercel UAT portals",
            "  UAT mobile (EAS/APK) ───┤  migrations · functions   ├── UAT mobile testers",
            "                         └── (dev + staging data)    │",
            "                                                       │",
            "                         ┌── OorjaMan Prod Supabase ─┤",
            "  GoDaddy prod web ───────┤                           ├── App Store / Play",
            "  Store mobile builds ────┤  live customer data       │",
            "                         └───────────────────────────┘",
        ],
        own_page=False,
    )
    add_table(
        doc,
        ["Surface", "UAT (today)", "Production (target)"],
        [
            ["customer-app", "EAS uat profile / local APK", "App Store + Google Play"],
            ["technician-app", "EAS uat profile / local APK", "App Store + Google Play"],
            ["admin-web", "Vercel", "GoDaddy admin.oorjaman.com"],
            ["vendor-web", "Vercel", "GoDaddy vendor.oorjaman.com"],
            ["support-web", "Vercel", "GoDaddy support.oorjaman.com"],
            ["oorjaman-web", "dev-oorjaman (planned)", "GoDaddy oorjaman.com"],
            ["Database", "OorjaMan UAT project", "OorjaMan Prod project"],
        ],
    )


def add_environments(doc: Document) -> None:
    doc.add_heading("16. Environments and configuration", level=1)
    add_table(
        doc,
        ["Tier", "Web env files", "Mobile env files", "Notes"],
        [
            ["Local dev", "apps/*/.env.development.local", "apps/*/.env.development.local", "Dummy auth typical"],
            ["UAT", "apps/*/.env.uat.local + Vercel Dashboard", "apps/*/.env.uat.local (embedded in APK)", "Vercel portals + UAT Supabase"],
            ["Production", "apps/*/.env.production.local", "apps/*/.env.production.local", "Real auth, Prod Supabase"],
        ],
    )
    doc.add_paragraph("Key variables (all apps): EXPO_PUBLIC_SUPABASE_URL, EXPO_PUBLIC_SUPABASE_ANON_KEY (mobile); VITE_SUPABASE_URL, VITE_SUPABASE_ANON_KEY (portals). Never ship service_role to clients.")


def add_external(doc: Document) -> None:
    doc.add_heading("17. External services", level=1)
    add_table(
        doc,
        ["Service", "Role", "Status"],
        [
            ["Supabase", "Database, auth, realtime, storage, edge functions", "Active (UAT); Prod planned"],
            ["Vercel", "UAT portal hosting", "Live"],
            ["Expo / EAS", "Mobile builds, push credentials", "UAT builds active"],
            ["GoDaddy", "Domain + marketing/portal hosting", "Planned for prod"],
            ["Google Maps", "Customer app maps, site photo geotag", "Optional until release"],
            ["Expo Push API", "Remote mobile notifications", "Via edge functions"],
            ["Apple / Google stores", "Production mobile distribution", "Not started"],
            ["Payment gateway", "—", "Not integrated (dummy payments)"],
        ],
    )


def add_security(doc: Document) -> None:
    doc.add_heading("18. Security overview", level=1)
    add_bullets(
        doc,
        [
            "All client access uses Supabase anon key + user JWT; RLS is the authorization boundary.",
            "service_role key is server/scripts only (seed, edge functions, CI secrets).",
            "Verification columns on public.users are trigger-guarded; only auth sync can set them.",
            "Push outbox tables have no client RLS access.",
            "Vercel portals: SPA with security headers in root vercel.json; see SECURITY-VERCEL.md.",
            "UAT marketing site uses noindex; production marketing is indexable only on oorjaman.com.",
        ],
    )


def add_appendix(doc: Document) -> None:
    doc.add_heading("19. Appendix", level=1)

    doc.add_heading("19.1 Root npm scripts", level=2)
    add_table(
        doc,
        ["Script", "Purpose"],
        [
            ["npm run customer / technician", "Start mobile Metro"],
            ["npm run admin / vendor / support", "Start portal dev servers"],
            ["npm run android:apk:uat:customer", "Build UAT Android APK"],
            ["npm run ios:uat:customer", "EAS iOS UAT build (requires Apple Developer)"],
            ["npm run db:push", "Apply Supabase migrations"],
            ["npm run functions:deploy", "Deploy edge functions"],
            ["npm run validate", "typecheck + knip"],
        ],
    )

    doc.add_heading("19.2 Related documentation", level=2)
    add_bullets(
        doc,
        [
            "project-docs/DEPLOYMENT.md — PROD vs UAT matrix",
            "project-docs/ENVIRONMENT.md — all environment variables",
            "project-docs/SUPABASE-UAT-PROD.md — dual-project workflow",
            "project-docs/VERCEL.md — portal hosting",
            "packages/api/src/database.types.ts — TypeScript schema reference",
            "supabase/schema.sql — generated DDL snapshot",
        ],
    )

    doc.add_heading("19.3 Document maintenance", level=2)
    doc.add_paragraph(
        f"Generated on {date.today().isoformat()} by scripts/generate-architecture-doc.py. "
        "Regenerate with: npm run docs:architecture"
    )


def build() -> None:
    doc = setup_document()
    add_cover(doc)
    add_toc_placeholder(doc)
    add_executive_summary(doc)
    add_master_architecture_diagram(doc)
    add_monorepo(doc)
    add_applications(doc)
    add_packages(doc)
    add_interactions(doc)
    add_auth(doc)
    add_flows(doc)
    add_notifications(doc)
    add_supabase_platform(doc)
    add_database(doc)
    add_rls(doc)
    add_edge_functions(doc)
    add_storage(doc)
    add_deployment(doc)
    add_environments(doc)
    add_external(doc)
    add_security(doc)
    add_appendix(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
